"""Read-only external P/Q audit, run with bundled document Python."""
import numpy as np
import pandas as pd
import openpyxl
from .contracts import EXTERNAL, EXTERNAL_SHA, OUT, PF
from .firewall import sha, write

def summary(values):
    a=np.asarray(values,dtype=float)
    a=a[np.isfinite(a)]
    if not len(a):
        return {'N':0}
    return {'N':len(a),'mean':float(a.mean()),'variance_population':float(a.var()),'min':float(a.min()),
            'P5':float(np.quantile(a,.05)),'P50':float(np.quantile(a,.5)),'P95':float(np.quantile(a,.95)),'max':float(a.max())}

def main():
    before=sha(EXTERNAL)
    assert before==EXTERNAL_SHA, 'EXTERNAL_SHA_MISMATCH'
    before_stat=EXTERNAL.stat()
    w=openpyxl.load_workbook(EXTERNAL,read_only=True,data_only=True)
    s=w.worksheets[0]
    it=s.iter_rows(values_only=True)
    headers=list(next(it))
    wanted=['timestamp','pa','pb','pc','ptotal','qa','qb','qc','qtotal','sa','sb','sc','stotal',
            'tpfa','tpfb','tpfc','tpftotal','vrmsa','vrmsb','vrmsc','irmsa','irmsb','irmsc']
    indices=[headers.index(c) for c in wanted]
    frame=pd.DataFrame([[row[i] for i in indices] for row in it],columns=wanted)
    w.close()
    for c in wanted:
        frame[c]=pd.to_numeric(frame[c],errors='coerce')
    p=frame[['pa','pb','pc']].sum(axis=1,min_count=3)
    q=frame[['qa','qb','qc']].sum(axis=1,min_count=3)
    apparent=frame[['sa','sb','sc']].sum(axis=1,min_count=3)
    valid=(apparent>0)&p.notna()&q.notna()
    pf=p.abs()/apparent.where(apparent>0)
    frame['PF_magnitude']=pf
    frame['P_consumption_W']=-p
    frame['Q_raw_var']=q
    frame['S_sum_VA']=apparent
    frame['Q_abs_over_P_abs']=q.abs()/p.abs().where(p.abs()>0)
    frame['time_utc']=pd.to_datetime(frame.timestamp,unit='ms',utc=True)
    local=frame.time_utc.dt.tz_convert('Europe/Madrid')
    byhour={str(h):summary(pf.loc[local.dt.hour==h]) for h in range(24)}
    bounds=np.quantile((-p)[valid],[1/3,2/3])
    levels=np.searchsorted(bounds,(-p).to_numpy(),side='right')
    load={name:{'P_W':summary((-p)[levels==i]),'PF':summary(pf[levels==i])} for i,name in enumerate(['low','medium','high'])}
    phase={c:summary(frame['p'+c].abs()/frame['s'+c].where(frame['s'+c]>0)) for c in 'abc'}
    checks={k:float(np.nanmax(np.abs(a-b))) for k,a,b in [
        ('total_P_minus_phase_sum_W',frame.ptotal,p),('total_Q_minus_phase_sum_var',frame.qtotal,q),
        ('total_S_minus_phase_sum_VA',frame.stotal,apparent),
        ('measured_tpf_percent_minus_100P_over_S',frame.tpftotal,100*frame.ptotal/frame.stotal)]}
    pfstats=summary(pf[valid])
    # External descriptive classification only, no runtime or AIDC parameter tuning.
    material=pfstats['P95']-pfstats['P5']>.02 or abs(pfstats['P50']-.95)>.02
    classification='MATERIAL_VARIABILITY_OBSERVED' if material else 'PLAUSIBLE_APPROXIMATION'
    report={'classification':classification,'role':'EXTERNAL_EMPIRICAL_REFERENCE_ONLY',
      'source':str(EXTERNAL),'source_url':'https://github.com/joaquinjgz/Dataset_university_data_center',
      'source_sha256_before':before,'source_sha256_after':sha(EXTERNAL),'SHA_PASS':sha(EXTERNAL)==before,
      'read_only':True,'mtime_unchanged':EXTERNAL.stat().st_mtime_ns==before_stat.st_mtime_ns,
      'worksheet':'Sheet1','source_rows_inclusive':[2,len(frame)+1],
      'columns_excel':{c:openpyxl.utils.get_column_letter(headers.index(c)+1) for c in wanted},
      'rows':len(frame),'valid_power_rows':int(valid.sum()),'invalid_rows':int((~valid).sum()),
      'time_range_UTC':[str(frame.time_utc.min()),str(frame.time_utc.max())],
      'duplicate_timestamps':int(frame.timestamp.duplicated().sum()),
      'interval_seconds':summary(np.diff(np.sort(frame.timestamp.to_numpy()))/1000),
      'P_raw_total_W':summary(p),'P_consumption_magnitude_W':summary(-p),'Q_raw_total_var':summary(q),'S_total_VA':summary(apparent),
      'PF_magnitude':pfstats,'PF_per_phase':phase,'Q_abs_over_P_abs':summary(frame.Q_abs_over_P_abs),
      'Q_raw_over_P_raw':summary(q/p.where(p!=0)),
      'PQ_equivalent_PF':summary(p.abs()/np.sqrt(p*p+q*q)),
      'definitions':{'P':'sum phase active W','Q':'sum phase reactive var','S':'sum measured phase apparent VA',
        'PF':'abs(sum P)/sum measured S; compare abs(tpftotal)/100; do not infer true PF from sqrt(P^2+Q^2) in harmonic data'},
      'polarity':'README documents inverted phase current sensors. P consumption sign reversed analytically; source unchanged. A simultaneous P,Q sign reversal would give Q_corrected=-Q_raw. Firmware Q convention is not fully specified, so leading/lagging remains CONDITIONAL_NOT_CERTIFIED.',
      'Q_positive_raw_fraction':float((q>0).mean()),'P_negative_raw_fraction':float((p<0).mean()),
      'leading_lagging':'Under signed complex-power reversal and passive-load convention: Q_corrected < 0 implies leading. Not asserted as unconditional facility authority.',
      'time_of_day_timezone':'Europe/Madrid','hourly_PF':byhour,'load_tertile_thresholds_W':bounds.tolist(),'PF_by_load':load,
      'load_PF_correlation':float(np.corrcoef((-p)[valid],pf[valid])[0,1]),
      'PF_095_absolute_error':summary(np.abs(pf-PF)),
      'fraction_abs_PF_error_gt_001':float((np.abs(pf-PF)>.01).mean()),
      'fixed_095_Q_magnitude_error_var':summary(abs(q)-abs(p)*np.tan(np.arccos(PF))),
      'classification_rule':'descriptive external screen: PF P95-P5 > .02 or median distance from .95 > .02 => material variability; otherwise plausible aggregate approximation, not per-phase fidelity',
      'reconciliation_max_abs':checks,'PF_gt_one_count':int((pf>1+1e-6).sum()),
      'AIDC_PF':PF,'AIDC_Q_CONTROL':'NO','facility_PQ_authority':'INDEPENDENT_LOCAL_DATA_REQUIRED',
      'UPS_STATCOM_capability_authority':'NOT_PROVIDED'}
    write('V40J_EXTERNAL_DATACENTER_PQ_AUDIT.json',report)
    frame[['time_utc','P_consumption_W','Q_raw_var','S_sum_VA','PF_magnitude','Q_abs_over_P_abs']].to_csv(OUT/'EXTERNAL_PQ_DERIVED_REFERENCE.csv',index=False,lineterminator='\n')
    text=f'''# 외부 데이터센터 P/Q 감사

판정: **{classification}**. University of Córdoba 외부 참고자료이며 우리 AIDC의 직접 authority가 아니다.

원본 SHA-256 `{before}`는 요구값과 일치하며 분석 전후 hash와 수정시각이 동일하다. Sheet1의 2–{len(frame)+1}행, {len(frame):,}개 측정값을 읽기 전용으로 분석했다.

측정 범위는 {report['time_range_UTC'][0]}부터 {report['time_range_UTC'][1]}까지다. 총 PF 크기의 P5/P50/P95는 {pfstats['P5']:.6f} / {pfstats['P50']:.6f} / {pfstats['P95']:.6f}, 분산은 {pfstats['variance_population']:.8f}다. 부하별·시간대별 분포와 상별 PF는 JSON에 저장했다.

총 P/Q는 상별 합, 총 S는 측정된 상별 apparent power 합을 사용했다. PF는 |P|/S다. 고조파 데이터에서 sqrt(P²+Q²)를 측정 S로 대체하지 않았다. 총합과 원본 total 열, PF와 원본 tpftotal(%)을 독립 대조했다.

[원자료 설명](https://github.com/joaquinjgz/Dataset_university_data_center)은 A/B/C 전류 센서 극성 반전을 명시한다. P 소비 방향만 분석값에서 정리했다. Q도 복소전력 부호 반전 대상이라는 가정에서는 leading으로 해석되나, 계측기 Q 부호 정의가 완전히 확인되지 않아 leading/lagging을 확정하지 않았다. 원본은 변경하지 않았다.

고정 PF=0.95의 외부 집계 근사 적합성과 상별·시간별 fidelity는 구분한다. 이번 revision은 AIDC PF=0.95와 Q control NO를 유지한다. 시변 외생 PF에는 우리 시설의 독립 P/Q 자료, 제어 Q에는 UPS/STATCOM의 P-Q capability authority가 필요하다.
'''
    (OUT/'V40J_EXTERNAL_DATACENTER_PQ_AUDIT.md').write_text(text,encoding='utf-8',newline='\n')
    print(classification, pfstats)

if __name__=='__main__':
    main()
